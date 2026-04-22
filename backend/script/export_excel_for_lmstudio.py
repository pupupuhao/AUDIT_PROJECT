from __future__ import annotations

from pathlib import Path
from typing import Any
from openpyxl import load_workbook


INPUT_XLSX = Path(
    "/home/jerrylmr/githubRepository/AUDIT_PROJECT/backend/modules/audit_engine/xlsx/辰弘佳苑小区联排及多层共13户屋面瓦片塌落、渗水维修.xlsx"
)
OUTPUT_TXT = Path(
    "/home/jerrylmr/githubRepository/AUDIT_PROJECT/backend/modules/audit_engine/xlsx/chenghongjia_yuan_for_lmstudio.txt"
)


def norm(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    return str(v).strip()


def row_is_empty(values: list[str]) -> bool:
    return all(v == "" for v in values)


def sheet_to_block(ws) -> str:
    rows: list[list[str]] = []
    max_col = ws.max_column
    max_row = ws.max_row

    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
        rows.append([norm(v) for v in row])

    while rows and row_is_empty(rows[-1]):
        rows.pop()

    merged_ranges = [str(r) for r in ws.merged_cells.ranges]

    lines: list[str] = []
    lines.append(f"===== Sheet: {ws.title} =====")
    lines.append(f"[meta] max_row={ws.max_row}, max_col={ws.max_column}")
    lines.append(
        f"[meta] merged_ranges={', '.join(merged_ranges)}" if merged_ranges else "[meta] merged_ranges=None"
    )

    lines.append("[raw_tsv]")
    for row in rows:
        lines.append("\t".join(row))

    lines.append("[cell_view]")
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            if value != "":
                col_letter = ws.cell(row=r_idx, column=c_idx).column_letter
                lines.append(f"{col_letter}{r_idx}={value}")

    return "\n".join(lines)


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"找不到文件: {INPUT_XLSX}")

    wb = load_workbook(INPUT_XLSX, data_only=True)

    preferred_order = [
        "维修工程信息",
        "维修对象",
        "维修预案",
        "维修决案",
        "业主表决汇总",
        "业主大会决议",
        "业主表决结果",
        "维修工单",
        "业主征询意见",
        "要件落款信息表",
    ]

    sheets = []
    existing = set(wb.sheetnames)
    for name in preferred_order:
        if name in existing:
            sheets.append(wb[name])
    for name in wb.sheetnames:
        if name not in preferred_order:
            sheets.append(wb[name])

    parts: list[str] = []
    parts.append("### Excel workbook exported for local LLM analysis ###")
    parts.append(f"workbook={INPUT_XLSX.name}")
    parts.append(f"sheet_count={len(wb.sheetnames)}")
    parts.append("sheet_names=" + ", ".join(wb.sheetnames))
    parts.append("")

    for ws in sheets:
        parts.append(sheet_to_block(ws))
        parts.append("")

    OUTPUT_TXT.write_text("\n".join(parts), encoding="utf-8")
    print(f"导出完成: {OUTPUT_TXT}")


if __name__ == "__main__":
    main()