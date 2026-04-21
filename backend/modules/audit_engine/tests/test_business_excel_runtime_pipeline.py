from io import BytesIO

from openpyxl import Workbook

from modules.audit_engine.services.excel_upload_service import parse_xlsx_bytes


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    workspace = workbook.create_sheet("维修工程信息")
    workspace.append(["工程id", "工程名称", "是否签定施工合同", "是否审价", "三审工程性质"])
    workspace.append(["WSID", "WSCODE", "WSNAME", "NEED_PRO_CONTRACT", "NEED_CHECK_AMT", "PROPERTY", "WS_AMT"])
    workspace.append(["1001", "1001", "1号楼屋面维修", "1", "1", "1", "50000"])

    draft = workbook.create_sheet("维修预案")
    draft.append(["预案ID", "工程名称", "是否签订施工合同", "是否审价", "预案金额"])
    draft.append(["BPPID", "WSCODE", "WSNAME", "PROPERTY", "NEED_PRO_CONTRACT", "NEED_CHECK_AMT", "ORGN_AMT", "WSID"])
    draft.append(["2001", "1001", "1号楼屋面维修", "1", "1", "1", "50000", "1001"])

    blueprint = workbook.create_sheet("维修决案")
    blueprint.append(["决案ID", "工程名称", "是否签订施工合同", "是否审价", "决案金额"])
    blueprint.append(["BPID", "WSCODE", "WSNAME", "PROPERTY", "NEED_PRO_CONTRACT", "NEED_CHECK_AMT", "ORGN_AMT", "FINAL_AMT", "WSID"])
    blueprint.append(["3001", "1001", "1号楼屋面维修", "1", "1", "1", "50000", "50000", "1001"])

    vote = workbook.create_sheet("业主表决汇总")
    vote.append(["BPPID", "COUNT_HOU", "SUM_AREA", "AGREE_HOU", "AGREE_AREA", "REQUEST_ENDDATE", "WSID"])
    vote.append(["2001", "10", "1000", "8", "800", "20240131", "1001"])

    project = workbook.create_sheet("三审工程维修项目表")
    project.append(["MO_ID", "BPID", "IS_SIGNED_ESC", "IS_SIGNED_ESR", "IS_SIGNED_PC", "NEED_CON", "ORGN_AMT", "FINAL_AMT", "STARTUP_DATE", "WSID"])
    project.append(["9001", "3001", "1", "1", "1", "1", "50000", "50000", "20240201", "1001"])

    contract = workbook.create_sheet("施工合同表")
    contract.append(["PCID", "NAME", "CONTRACT_AMT", "STARTUP_DATE", "SIGN_DATE", "WSID"])
    contract.append(["4001", "1号楼屋面维修合同", "49000", "20240201", "20240120", "1001"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_business_excel_outputs_runtime_standard_fields_from_multiple_sheets():
    parsed = parse_xlsx_bytes(_workbook_bytes(), filename="demo.xlsx")

    assert parsed["parse_mode"] == "business_package"
    assert len(parsed["rows"]) == 1
    row = parsed["rows"][0]
    fields = row["standard_fields"]
    assert row["project_key"] == "1001"
    assert fields["project_name"]["value"] == "1号楼屋面维修"
    assert fields["need_cost_review"]["value"] is True
    assert fields["has_appraisal_contract"]["value"] is True
    assert fields["has_appraisal_report"]["value"] is True
    assert fields["has_construction_contract"]["value"] is True
    assert fields["budget_amount"]["value"] == 50000.0
    assert fields["contract_amount"]["value"] == 49000.0
    assert fields["vote_date"]["value"] == "2024-01-31"
    assert fields["vote_legal"]["value"] is True
    assert row["debug"]["unmapped_columns"] == []


def test_conflicting_source_field_is_exposed_without_blocking_value_selection():
    parsed = parse_xlsx_bytes(_workbook_bytes(), filename="demo.xlsx")
    row = parsed["rows"][0]

    assert row["standard_fields"]["need_cost_review"]["status"] == "resolved"
    assert row["standard_fields"]["need_cost_review"]["value"] is True

